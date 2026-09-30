"""Round 5.1: an MRV export names what its Carbon number is.

Round 5 could not show the factor version inside an export: the per-season
Carbon block carried only a factor-set UUID, and the PDF printed that UUID, the
raw scenario `actual` and a season UUID as heading. Each season's Carbon block
now also states the KIND of result (always the actual one — the export never
reads a scenario) and the factor-set VERSION, taken from the same factor-set
row the provenance section reads; an unreadable set is `null`, never guessed.
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

from pypdf import PdfReader

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from mrv import report_pdf  # noqa: E402
from tests.test_mrv_export import FACTOR_SET, FACTOR_SET_ROW, SEASON, SUCCEEDED_CARBON, FakeCarbon, FakeRead, generate  # noqa: E402


_ACTUAL = {**SUCCEEDED_CARBON, "scenario": "actual"}


def _manifest(factor_sets=FACTOR_SET_ROW, carbon=_ACTUAL):
    result, _, _, _ = generate(read=FakeRead(factor_sets=factor_sets), carbon=FakeCarbon(carbon))
    return result["manifest"]


def test_each_season_names_the_actual_result_and_its_factor_version():
    c = _manifest()["carbon"]["per_crop_season"][SEASON]
    assert c["calculation_kind"] == "actual"
    assert c["ef_config_version"] == FACTOR_SET_ROW[FACTOR_SET]["version_code"]
    assert c["factor_set_id"] == FACTOR_SET


def test_an_unreadable_factor_set_gives_no_invented_version():
    c = _manifest(factor_sets={})["carbon"]["per_crop_season"][SEASON]
    assert c["factor_set_id"] == FACTOR_SET and c["ef_config_version"] is None


def test_a_season_without_a_result_says_so_without_kind_or_version():
    manifest = _manifest(carbon=None)
    c = manifest["carbon"]["per_crop_season"][SEASON]
    assert c["status"] == "unavailable" and c["calculation_kind"] is None and c["ef_config_version"] is None


def test_pdf_labels_kind_scenario_version_and_status_instead_of_raw_values():
    manifest = _manifest()
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(report_pdf.render_pdf(manifest))).pages)
    flat = " ".join(text.split())
    assert "Kết quả vận hành" in flat and "Theo dữ liệu đã ghi nhận" in flat
    # People read labels; codes stay in the JSON/XLSX data (Round 5.1).
    for raw in ("(actual)", "(cooperative_manager)", "(draft)", "irrigation_ch4", "fertilizer_n2o"):
        assert raw not in flat
    assert FACTOR_SET_ROW[FACTOR_SET]["version_code"] in flat
    assert "undefined" not in flat.lower()
    # The factor set is named by its version in the Carbon block, not by UUID.
    carbon_part = flat.split("Loại kết quả", 1)[1].split("Nguồn gốc hệ số", 1)[0]
    assert FACTOR_SET not in carbon_part


def test_workbook_carbon_sheet_carries_kind_and_version():
    from openpyxl import load_workbook
    from mrv import workbook

    wb = load_workbook(io.BytesIO(workbook.render_workbook(_manifest())))
    ws = wb["Carbon"]
    header = next(r for r in ws.iter_rows(values_only=True) if r and r[0] == "Mã vụ")
    row = next(r for r in ws.iter_rows(values_only=True) if r and r[0] == SEASON)
    values = dict(zip(header, row))
    assert values["Loại kết quả"] == "actual"
    assert values["Phiên bản bộ hệ số"] == FACTOR_SET_ROW[FACTOR_SET]["version_code"]
