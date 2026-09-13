"""M07 part 2 — the XLSX renderer and its artifacts.

The contract under test is narrow and strict: a workbook is a *rendering* of one
stored manifest. It must contain what the snapshot says, nothing newer, nothing
recomputed, and it must keep a null a null.

Workbooks are written to bytes and read back with openpyxl rather than compared
as files: a .xlsx is a zip, and zip metadata is not stable, so byte-comparing
would test the archiver rather than the report.
"""
from __future__ import annotations

import copy
import hashlib
import io
import json
import math
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import api  # noqa: E402
from mrv import manifest as m  # noqa: E402
from mrv import workbook as wb  # noqa: E402
from service import MrvExportService  # noqa: E402
from tests.test_mrv_export import (  # noqa: E402
    ACTOR, CASE, FACTOR_SET_ROW, ORG, SEASON, SUCCEEDED_CARBON,
    FakeCarbon, FakeExportStore, FakeRead,
)

RENDERED_AT = datetime(2026, 9, 13, 9, 0, 0, tzinfo=timezone.utc)


def build_manifest(read=None, carbon=None):
    store = FakeExportStore()
    service = MrvExportService(store, carbon or FakeCarbon())
    repo = read or FakeRead()
    created = service.create(read_repository=repo, mrv_case_id=CASE)
    return created["manifest"], service, repo, store


def render(manifest) -> dict:
    """Render and read back, returning {sheet_title: worksheet}."""
    data = wb.render_workbook(manifest, rendered_at=RENDERED_AT)
    book = load_workbook(io.BytesIO(data))
    return {name: book[name] for name in book.sheetnames}


def cells(ws, col: int) -> list:
    return [ws.cell(row=r, column=col).value for r in range(1, ws.max_row + 1)]


def find_row(ws, label: str, col: int = 1) -> int | None:
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=col).value == label:
            return r
    return None


def header_row(ws, first_header: str, col: int = 1) -> int:
    row = find_row(ws, first_header, col)
    assert row, f"header {first_header!r} not found"
    return row


# --------------------------------------------------------------------------
# Structure
# --------------------------------------------------------------------------

def test_workbook_has_the_documented_sheets_in_order():
    manifest, *_ = build_manifest()
    data = wb.render_workbook(manifest, rendered_at=RENDERED_AT)
    book = load_workbook(io.BytesIO(data))
    assert book.sheetnames == wb.SHEET_TITLES
    # The default openpyxl sheet must not survive.
    assert "Sheet" not in book.sheetnames


def test_every_sheet_is_present_even_when_its_section_is_empty():
    """A section with no data gets an explicit empty-state row, not a silent gap."""
    empty = FakeRead(evidence=[], scope=False, activities=[])
    manifest, *_ = build_manifest(read=empty)
    sheets = render(manifest)
    assert set(sheets) == set(wb.SHEET_TITLES)
    for title in ("Bằng chứng", "Phạm vi", "Hoạt động canh tác", "Thu hoạch"):
        flat = [v for v in cells(sheets[title], 1) if v]
        assert wb.EMPTY_NOTE in flat, f"{title} lacks an empty-state row"


# --------------------------------------------------------------------------
# Determinism
# --------------------------------------------------------------------------

def test_same_manifest_renders_identical_content_twice():
    manifest, *_ = build_manifest()
    first, second = render(manifest), render(manifest)
    for title in wb.SHEET_TITLES:
        a, b = first[title], second[title]
        assert a.max_row == b.max_row and a.max_column == b.max_column
        assert [cells(a, c) for c in range(1, a.max_column + 1)] == \
               [cells(b, c) for c in range(1, b.max_column + 1)]


def test_row_order_does_not_follow_manifest_list_order():
    """Shuffling the manifest's lists must not change the workbook."""
    manifest, *_ = build_manifest()
    shuffled = copy.deepcopy(manifest)
    for key in ("steps", "evidence", "activities", "warnings"):
        shuffled[key] = list(reversed(shuffled.get(key) or []))
    shuffled["harvest"]["events"] = list(reversed(shuffled["harvest"]["events"]))

    a, b = render(manifest), render(shuffled)
    for title in ("Các bước MRV", "Bằng chứng", "Hoạt động canh tác", "Cảnh báo", "Thu hoạch"):
        assert [cells(a[title], c) for c in range(1, a[title].max_column + 1)] == \
               [cells(b[title], c) for c in range(1, b[title].max_column + 1)], title


# --------------------------------------------------------------------------
# Null, type and date semantics
# --------------------------------------------------------------------------

def test_a_null_metric_is_a_blank_cell_not_zero():
    manifest, *_ = build_manifest()
    ws = render(manifest)["Chỉ số tài nguyên"]
    head = header_row(ws, "Mã vụ")
    labels = cells(ws, 2)
    values = cells(ws, 3)
    water_row = labels.index("Nước trên mỗi kg")
    assert values[water_row] is None, "a missing metric must be an empty cell"
    fert_row = labels.index("Phân bón trên mỗi kg")
    assert values[fert_row] == 0.02317
    del head


def test_no_cell_anywhere_says_null_or_none_or_na():
    manifest, *_ = build_manifest()
    for ws in render(manifest).values():
        for row in ws.iter_rows(values_only=True):
            for value in row:
                if isinstance(value, str):
                    assert value.strip().lower() not in ("null", "none", "n/a", "nan")


def test_numbers_are_numeric_cells_with_precision_preserved():
    manifest, *_ = build_manifest()
    ws = render(manifest)["Thu hoạch"]
    head = header_row(ws, "Mã hoạt động")
    row = head + 1
    yield_kg = ws.cell(row=row, column=4).value
    area = ws.cell(row=row, column=5).value
    assert isinstance(yield_kg, (int, float)) and not isinstance(yield_kg, bool)
    assert yield_kg == 5200
    assert area == 1.5


def test_datetimes_are_real_datetime_cells_in_utc():
    manifest, *_ = build_manifest()
    ws = render(manifest)["Tổng quan"]
    row = find_row(ws, "Thời điểm tạo (UTC)")
    value = ws.cell(row=row, column=2).value
    assert isinstance(value, datetime)
    assert value.tzinfo is None, "Excel cells carry no timezone"
    # The instant is preserved as UTC, not shifted into local time.
    assert value == datetime.fromisoformat(
        manifest["generated_at"].replace("Z", "+00:00")
    ).replace(tzinfo=None)


def test_dates_stay_dates_and_do_not_become_midnight_timestamps():
    from datetime import date

    manifest, *_ = build_manifest()
    ws = render(manifest)["Tổng quan"]
    row = find_row(ws, "Kỳ bắt đầu")
    value = ws.cell(row=row, column=2).value
    assert isinstance(value, (date, datetime))
    assert (value.date() if isinstance(value, datetime) else value) == date(2026, 5, 1)


def test_no_formulas_are_written():
    """The workbook is a report, not a second calculation engine."""
    manifest, *_ = build_manifest()
    for ws in render(manifest).values():
        for row in ws.iter_rows(values_only=True):
            for value in row:
                assert not (isinstance(value, str) and value.startswith("=")), value


# --------------------------------------------------------------------------
# Content fidelity: the snapshot, nothing else
# --------------------------------------------------------------------------

def test_step_statuses_are_preserved_verbatim():
    manifest, *_ = build_manifest()
    ws = render(manifest)["Các bước MRV"]
    head = header_row(ws, "Thứ tự")
    statuses = [ws.cell(row=r, column=3).value for r in range(head + 1, ws.max_row + 1)]
    assert statuses == [s["status"] for s in sorted(manifest["steps"], key=lambda x: x["step_no"])]
    assert "failed" not in statuses


def test_evidence_is_metadata_only_and_says_so():
    manifest, *_ = build_manifest()
    ws = render(manifest)["Bằng chứng"]
    note = " ".join(str(v) for v in cells(ws, 1) if v)
    assert "KHÔNG kèm" in note
    head = header_row(ws, "Mã bằng chứng")
    included = ws.cell(row=head + 1, column=10).value
    assert included == "không"


def test_carbon_unavailable_renders_a_status_not_a_zero():
    manifest, *_ = build_manifest()
    ws = render(manifest)["Carbon"]
    head = header_row(ws, "Mã vụ")
    assert ws.cell(row=head + 1, column=2).value == "unavailable"
    assert ws.cell(row=head + 1, column=3).value == "no_succeeded_calculation"
    assert ws.cell(row=head + 1, column=7).value is None, "total CO2e must be blank, not 0"


def test_carbon_available_renders_values_and_breakdown():
    read = FakeRead(factor_sets=copy.deepcopy(FACTOR_SET_ROW))
    manifest, *_ = build_manifest(read=read, carbon=FakeCarbon(SUCCEEDED_CARBON))
    ws = render(manifest)["Carbon"]
    head = header_row(ws, "Mã vụ")
    assert ws.cell(row=head + 1, column=2).value == "succeeded"
    assert ws.cell(row=head + 1, column=7).value == 2920.8
    breakdown_head = header_row(ws, "Hạng mục", col=2)
    assert ws.cell(row=breakdown_head + 1, column=2).value == "methane"
    assert ws.cell(row=breakdown_head + 1, column=8).value == 2912


def test_provenance_unavailable_is_shown_not_hidden():
    manifest, *_ = build_manifest()
    ws = render(manifest)["Nguồn gốc hệ số"]
    flat = [str(v) for v in cells(ws, 1) if v]
    assert "factor_provenance_unavailable" in flat


def test_provenance_renders_real_sources_without_inventing_any():
    read = FakeRead(factor_sets=copy.deepcopy(FACTOR_SET_ROW))
    manifest, *_ = build_manifest(read=read, carbon=FakeCarbon(SUCCEEDED_CARBON))
    ws = render(manifest)["Nguồn gốc hệ số"]
    text = " ".join(str(v) for row in ws.iter_rows(values_only=True) for v in row if v)
    assert "IPCC 2019 Refinement" in text
    assert "Table 5.11" in text
    assert "factor_provenance_unavailable" not in text


def test_warnings_are_one_row_each_with_severity_unchanged():
    manifest, *_ = build_manifest()
    ws = render(manifest)["Cảnh báo"]
    head = header_row(ws, "Mã cảnh báo")
    rendered = [
        (ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value)
        for r in range(head + 1, ws.max_row + 1)
    ]
    expected = sorted(
        (w["code"], w["severity"]) for w in manifest["warnings"]
    )
    assert sorted(rendered) == expected
    assert all(sev in ("info", "warning") for _, sev in rendered)


def test_soft_deleted_activities_stay_out_and_the_sheet_says_so():
    manifest, *_ = build_manifest()
    ws = render(manifest)["Hoạt động canh tác"]
    head = header_row(ws, "Mã hoạt động")
    ids = [ws.cell(row=r, column=1).value for r in range(head + 1, ws.max_row + 1)]
    assert ids == ["a1"]
    note = " ".join(str(v) for v in cells(ws, 1)[:head] if v)
    assert "KHÔNG được" in note


def test_activity_detail_columns_do_not_bleed_between_types():
    manifest, *_ = build_manifest()
    ws = render(manifest)["Hoạt động canh tác"]
    head = header_row(ws, "Mã hoạt động")
    headers = [ws.cell(row=head, column=c).value for c in range(1, ws.max_column + 1)]
    row = head + 1  # the single harvest activity
    water_col = headers.index("Lượng nước (m³)") + 1
    yield_col = headers.index("Sản lượng (kg)") + 1
    assert ws.cell(row=row, column=yield_col).value == 5200
    assert ws.cell(row=row, column=water_col).value is None


def test_manifest_sheet_carries_the_integrity_block():
    manifest, *_ = build_manifest()
    ws = render(manifest)["Gói dữ liệu gốc"]
    sha_row = find_row(ws, "Mã băm gói dữ liệu (manifest_sha256)")
    assert ws.cell(row=sha_row, column=2).value == \
        manifest["package_integrity"]["manifest_sha256"]
    assert ws.cell(row=find_row(ws, "Phạm vi băm"), column=2).value == \
        "manifest-without-package_integrity"
    assert ws.cell(row=find_row(ws, "Bảng tính kết xuất lúc (UTC)"), column=2).value == \
        RENDERED_AT.replace(tzinfo=None)


def test_no_certification_language_in_the_workbook():
    read = FakeRead(factor_sets=copy.deepcopy(FACTOR_SET_ROW))
    manifest, *_ = build_manifest(read=read, carbon=FakeCarbon(SUCCEEDED_CARBON))
    text = " ".join(
        str(v).lower()
        for ws in render(manifest).values()
        for row in ws.iter_rows(values_only=True)
        for v in row if v
    )
    for claim in ("certified", "mrv compliant", "mrv_compliant", "verified by",
                  "chứng nhận đạt", "đã thẩm định", "official"):
        assert claim not in text, claim
    assert "không phải chứng nhận" in text


def test_the_renderer_refuses_something_that_is_not_a_manifest():
    with pytest.raises(ValueError):
        wb.render_workbook({"nope": 1})


def test_filename_is_deterministic_and_ascii():
    name = wb.workbook_filename("Hồ sơ/2026", "2026-09-13T08:00:00Z",
                                "abcdef01-2345-6789-abcd-ef0123456789")
    assert name == "agricarbon-mrv-Ho-so-2026-2026-09-13-abcdef01.xlsx"
    assert name.isascii() and "/" not in name


# --------------------------------------------------------------------------
# Service: lineage, integrity, immutability
# --------------------------------------------------------------------------

def test_xlsx_export_records_its_source_snapshot():
    store = FakeExportStore()
    service = MrvExportService(store, FakeCarbon())
    repo = FakeRead()
    snapshot = service.create(read_repository=repo, mrv_case_id=CASE, fmt="json")
    artifact = service.render(read_repository=repo, export_id=snapshot["export_id"])

    assert artifact["format"] == "xlsx"
    assert artifact["source_snapshot_export_id"] == snapshot["export_id"]
    assert artifact["export_id"] != snapshot["export_id"]
    # Same data, different file: that is exactly what the two digests say.
    assert artifact["payload_sha256"] == snapshot["payload_sha256"]
    assert artifact["file_sha256"] != snapshot["file_sha256"]
    assert store.rows[artifact["export_id"]]["export_payload"] == snapshot["manifest"]


def test_creating_an_xlsx_directly_still_produces_one_snapshot_and_one_artifact():
    store = FakeExportStore()
    service = MrvExportService(store, FakeCarbon())
    repo = FakeRead()
    artifact = service.create(read_repository=repo, mrv_case_id=CASE, fmt="xlsx")

    rows = list(store.rows.values())
    assert sorted(r["format"] for r in rows) == ["json", "xlsx"]
    snapshot_id = artifact["source_snapshot_export_id"]
    assert store.rows[snapshot_id]["format"] == "json"


def test_file_sha256_matches_the_stored_object():
    store = FakeExportStore()
    service = MrvExportService(store, FakeCarbon())
    repo = FakeRead()
    artifact = service.create(read_repository=repo, mrv_case_id=CASE, fmt="xlsx")
    row = store.rows[artifact["export_id"]]
    stored = store.objects[row["storage_object_path"]]
    assert hashlib.sha256(stored).hexdigest() == row["file_sha256"]
    assert artifact["byte_size"] == len(stored)


def test_xlsx_renders_the_old_snapshot_after_source_data_changes():
    """The critical one. A workbook must report the snapshot, not the present."""
    store = FakeExportStore()
    service = MrvExportService(store, FakeCarbon())
    repo = FakeRead()
    snapshot_a = service.create(read_repository=repo, mrv_case_id=CASE, fmt="json")

    repo._evidence.append({
        "id": "e-late", "step_no": 2, "production_batch_id": None, "evidence_type": "doc",
        "file_name": "late.pdf", "mime_type": "application/pdf", "sha256": "b" * 64,
        "storage_bucket": "mrv-evidence", "storage_object_path": f"{ORG}/{CASE}/late.pdf",
        "uploaded_by": ACTOR, "uploaded_at": datetime(2026, 9, 14, tzinfo=timezone.utc),
    })

    artifact_a = service.render(read_repository=repo, export_id=snapshot_a["export_id"])
    data, _, _ = service.download(read_repository=repo, export_id=artifact_a["export_id"])
    ws = load_workbook(io.BytesIO(data))["Bằng chứng"]
    head = header_row(ws, "Mã bằng chứng")
    ids = [ws.cell(row=r, column=1).value for r in range(head + 1, ws.max_row + 1)]
    assert ids == ["e1"], "the workbook must not contain evidence added after the snapshot"

    snapshot_b = service.create(read_repository=repo, mrv_case_id=CASE, fmt="json")
    artifact_b = service.render(read_repository=repo, export_id=snapshot_b["export_id"])
    data_b, _, _ = service.download(read_repository=repo, export_id=artifact_b["export_id"])
    ws_b = load_workbook(io.BytesIO(data_b))["Bằng chứng"]
    head_b = header_row(ws_b, "Mã bằng chứng")
    ids_b = [ws_b.cell(row=r, column=1).value for r in range(head_b + 1, ws_b.max_row + 1)]
    assert ids_b == ["e1", "e-late"]


def test_rendering_from_a_rendering_is_refused():
    store = FakeExportStore()
    service = MrvExportService(store, FakeCarbon())
    repo = FakeRead()
    artifact = service.create(read_repository=repo, mrv_case_id=CASE, fmt="xlsx")
    with pytest.raises(Exception) as exc:
        service.render(read_repository=repo, export_id=artifact["export_id"])
    assert exc.type.__name__ == "UnsupportedExportFormatError"


def test_a_missing_object_is_a_controlled_failure_not_a_rebuild():
    store = FakeExportStore()
    service = MrvExportService(store, FakeCarbon())
    repo = FakeRead()
    artifact = service.create(read_repository=repo, mrv_case_id=CASE, fmt="xlsx")
    store.objects.clear()  # object lost; metadata intact
    with pytest.raises(Exception) as exc:
        service.download(read_repository=repo, export_id=artifact["export_id"])
    assert exc.type.__name__ == "MrvArtifactMissingError"


def test_tampered_bytes_fail_closed():
    store = FakeExportStore()
    service = MrvExportService(store, FakeCarbon())
    repo = FakeRead()
    artifact = service.create(read_repository=repo, mrv_case_id=CASE, fmt="xlsx")
    path = store.rows[artifact["export_id"]]["storage_object_path"]
    store.objects[path] = store.objects[path] + b"tampered"
    with pytest.raises(Exception) as exc:
        service.download(read_repository=repo, export_id=artifact["export_id"])
    assert exc.type.__name__ == "MrvArtifactCorruptError"


# --------------------------------------------------------------------------
# HTTP
# --------------------------------------------------------------------------

@pytest.fixture
def client_for():
    def build(read, carbon=None, store=None):
        app = FastAPI()
        app.include_router(api.router)
        store = store or FakeExportStore()
        service = MrvExportService(store, carbon or FakeCarbon())
        app.dependency_overrides[api._read_repo] = lambda: read
        app.dependency_overrides[api._mrv_export_service] = lambda: service
        return TestClient(app), store
    return build


def test_manager_generates_and_downloads_xlsx(client_for):
    client, _ = client_for(FakeRead())
    created = client.post(f"/v1/mrv/cases/{CASE}/exports", json={"format": "xlsx"})
    assert created.status_code == 201
    body = created.json()
    assert body["format"] == "xlsx"
    assert body["file_name"].endswith(".xlsx")

    response = client.get(f"/v1/mrv/exports/{body['export_id']}/download")
    assert response.status_code == 200
    assert response.headers["content-type"] == \
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment;" in response.headers["content-disposition"]
    assert ".xlsx" in response.headers["content-disposition"]
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == wb.SHEET_TITLES


def test_the_response_never_names_the_bucket_or_object_path(client_for):
    client, _ = client_for(FakeRead())
    body = client.post(f"/v1/mrv/cases/{CASE}/exports", json={"format": "xlsx"}).json()
    text = json.dumps(body)
    assert "storage_object_path" not in text and "storage_bucket" not in text
    assert "mrv-exports" not in text
    assert body["file_name"] and "/" not in body["file_name"]


def test_a_farmer_cannot_generate_xlsx(client_for):
    client, _ = client_for(FakeRead(roles=("farmer",)))
    assert client.post(f"/v1/mrv/cases/{CASE}/exports",
                       json={"format": "xlsx"}).status_code == 404


def test_a_farmer_cannot_download_a_known_xlsx_export(client_for):
    manager, store = client_for(FakeRead())
    created = manager.post(f"/v1/mrv/cases/{CASE}/exports", json={"format": "xlsx"}).json()
    farmer, _ = client_for(FakeRead(roles=("farmer",)), store=store)
    assert farmer.get(f"/v1/mrv/exports/{created['export_id']}/download").status_code == 404


def test_a_farmer_cannot_render_from_a_known_snapshot(client_for):
    manager, store = client_for(FakeRead())
    snapshot = manager.post(f"/v1/mrv/cases/{CASE}/exports").json()
    farmer, _ = client_for(FakeRead(roles=("farmer",)), store=store)
    assert farmer.post(f"/v1/mrv/exports/{snapshot['export_id']}/render").status_code == 404


def test_a_cross_org_manager_cannot_generate_or_download_xlsx(client_for):
    manager, store = client_for(FakeRead())
    created = manager.post(f"/v1/mrv/cases/{CASE}/exports", json={"format": "xlsx"}).json()
    stranger = FakeRead(roles=("cooperative_manager",),
                        membership_org="99999999-9999-9999-9999-999999999999")
    other, _ = client_for(stranger, store=store)
    assert other.post(f"/v1/mrv/cases/{CASE}/exports", json={"format": "xlsx"}).status_code == 404
    assert other.get(f"/v1/mrv/exports/{created['export_id']}/download").status_code == 404


def test_unauthenticated_xlsx_is_401():
    app = FastAPI()
    app.include_router(api.router)
    client = TestClient(app)
    assert client.post(f"/v1/mrv/cases/{CASE}/exports", json={"format": "xlsx"}).status_code == 401
    assert client.post(f"/v1/mrv/exports/{CASE}/render").status_code == 401


def test_an_unknown_format_is_refused_with_the_shared_envelope(client_for):
    client, _ = client_for(FakeRead())
    response = client.post(f"/v1/mrv/cases/{CASE}/exports", json={"format": "docx"})
    assert response.status_code == 422
    assert response.json()["detail"]["error"]["code"] == "unsupported_export_format"


def test_a_missing_object_returns_404_with_the_shared_envelope(client_for):
    client, store = client_for(FakeRead())
    created = client.post(f"/v1/mrv/cases/{CASE}/exports", json={"format": "xlsx"}).json()
    store.objects.clear()
    response = client.get(f"/v1/mrv/exports/{created['export_id']}/download")
    assert response.status_code == 404
    assert response.json()["detail"]["error"]["code"] == "export_artifact_missing"


def test_a_tampered_object_returns_409_and_no_partial_body(client_for):
    client, store = client_for(FakeRead())
    created = client.post(f"/v1/mrv/cases/{CASE}/exports", json={"format": "xlsx"}).json()
    path = store.rows[created["export_id"]]["storage_object_path"]
    store.objects[path] = b"not a workbook"
    response = client.get(f"/v1/mrv/exports/{created['export_id']}/download")
    assert response.status_code == 409
    assert response.json()["detail"]["error"]["code"] == "export_artifact_integrity_failed"


# --------------------------------------------------------------------------
# Scale
# --------------------------------------------------------------------------

def _large_manifest(base: dict, *, activities: int, evidence: int) -> dict:
    big = copy.deepcopy(base)
    template = big["activities"][0]
    big["activities"] = []
    for i in range(activities):
        row = copy.deepcopy(template)
        row["activity_id"] = f"act-{i:05d}"
        row["occurred_at"] = f"2026-0{(i % 9) + 1}-01T03:00:00Z"
        big["activities"].append(row)
    ev = big["evidence"][0]
    big["evidence"] = []
    for i in range(evidence):
        row = copy.deepcopy(ev)
        row["evidence_id"] = f"ev-{i:05d}"
        row["step_no"] = (i % 6) + 1
        big["evidence"].append(row)
    return big


def test_a_large_manifest_renders_in_reasonable_time_and_size():
    manifest, *_ = build_manifest()
    big = _large_manifest(manifest, activities=1000, evidence=100)

    started = time.monotonic()
    data = wb.render_workbook(big, rendered_at=RENDERED_AT)
    elapsed = time.monotonic() - started

    book = load_workbook(io.BytesIO(data))
    acts = book["Hoạt động canh tác"]
    head = header_row(acts, "Mã hoạt động")
    assert acts.max_row - head == 1000
    assert book["Bằng chứng"].max_row - header_row(book["Bằng chứng"], "Mã bằng chứng") == 100

    # Generous: this guards against an accidental quadratic, not a few hundred ms.
    assert elapsed < 30, f"1000-activity render took {elapsed:.1f}s"
    assert len(data) < 8 * 1024 * 1024, f"workbook is {len(data)} bytes"
    print(f"\n  large manifest: {elapsed:.2f}s, {len(data) / 1024:.0f} KiB")


def test_numbers_keep_every_digit_excel_can_hold():
    # Found by the hosted smoke: openpyxl serializes numbers as "%.16g", so a
    # 17-digit ratio loses its last digit. Pin exactly that, and nothing worse:
    # the cell stays numeric and agrees to 16 significant digits.
    manifest, *_ = build_manifest()
    raw = "0.028846153846153848"
    manifest["resource_metrics"] = {"per_crop_season": {"season-x": {"fertilizer_per_kg": raw}}}
    book = load_workbook(io.BytesIO(wb.render_workbook(manifest, rendered_at=RENDERED_AT)))
    cells = [row[2] for row in book["Chỉ số tài nguyên"].iter_rows(values_only=True)
             if row[1] == "Phân bón trên mỗi kg"]
    assert len(cells) == 1 and isinstance(cells[0], float)
    assert cells[0] == float("%.16g" % float(raw))
    assert math.isclose(cells[0], float(raw), rel_tol=1e-15, abs_tol=0)
